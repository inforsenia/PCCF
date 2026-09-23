#!/usr/bin/env python3
"""Prepara la còpia de muntatge (STAGE) de les PDs abans de compilar el PDF
de Programacions (`compila-pd-pccf-%`). Mai toca els fitxers dels docents.

Per a cada PD de mòdul copiada a STAGE:
  1. Marca el títol de nivel superior (`# ...`, i per tant l'índex) amb
     ✎ si està en BORRADOR i amb ✗ si té incidències pendents (placeholders
     `[###]`/`[...]` o Excel incoherent per a la seua fulla) -- mateix criteri
     visual que ✏️/❌ a memòries, però amb pifont perquè el PD compila amb
     xelatex (el paquet emoji requereix LuaTeX).
  2. Elimina els blocs de cita (`> ...`): a les plantilles de PD només
     s'usen per a instruccions al docent, que no han d'eixir al PDF (mateix
     regex que memòries, `memories_utils.py`).
  3. Afig al final el Quadre Resum de la fulla del mòdul a l'Excel
     (`libro_{CICLE}.xlsx` de PD_DIR, el que editen els docents), exportat a
     PDF amb LibreOffice i inclòs amb \\includepdf. El títol
     `## Esquema general de ...` de la plantilla es trau del markdown i
     s'afig a l'índex amb `addtotoc` del mateix \\includepdf: si no, quedava
     sol en una pàgina en blanc abans del Quadre Resum (apaïsat).

A més, escriu STAGE/.draft si el report no està verificat
(`is_pd_verified`), perquè el Makefile active la marca d'aigua ESBORRANY.
"""

import argparse
import importlib.util
import json
import os
import re
import shutil
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pccf_utils import parse_pd_filename, get_hoja_label
from report_pccf import compute_pd_status, is_pd_verified, find_placeholders

PROJECT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

MARCA_BORRADOR = r"\texorpdfstring{\ \ding{46}}{}"
MARCA_INCIDENCIES = r"\texorpdfstring{\ \textcolor{red}{\ding{55}}}{}"
HEADER_TEX = "\\usepackage{pifont}\n"
BLOCKQUOTE_RE = re.compile(r'(?:^|\n)[ \t]*>.*(?:\n[ \t]*>.*)*')


def load_excel_exporter():
    """excel-to-pdfs.py té guionet al nom: no es pot importar amb `import`."""
    path = os.path.join(PROJECT_DIR, "tools", "excel-to-pdfs.py")
    spec = importlib.util.spec_from_file_location("excel_to_pdfs", path)
    src = open(path, encoding="utf-8").read()
    # Només volem la funció, no el codi de nivell de mòdul (llig sys.argv).
    src = src.split("\nruta_excel = sys.argv[1]")[0]
    mod = importlib.util.module_from_spec(spec)
    exec(compile(src, path, "exec"), mod.__dict__)
    return mod.exportar_rango_a_pdf


def resol_fulla(sheetnames, nombre):
    """json2excel crea la fulla amb el nom complet del mòdul; els llibres
    antics/proporcionats usen les sigles (get_hoja_label). Excel talla a 31."""
    for cand in (nombre, nombre[:31], get_hoja_label(nombre)):
        if cand in sheetnames:
            return cand
    return None


def marca_titol(path, marca):
    with open(path, encoding="utf-8") as f:
        lines = f.read().split("\n")
    for i, line in enumerate(lines):
        if line.startswith("# "):
            lines[i] = line.rstrip() + marca
            break
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))


ESQUEMA_RE = re.compile(r'^## (Esquema general de .*?)\s*$', re.MULTILINE)
LATEX_ESPECIALS = {"&": r"\&", "%": r"\%", "#": r"\#", "_": r"\_", "$": r"\$"}


def extrau_titol_esquema(path):
    """Lleva la línia `## Esquema general de ...` i en torna el text (o None)."""
    with open(path, encoding="utf-8") as f:
        content = f.read()
    m = ESQUEMA_RE.search(content)
    if not m:
        return None
    with open(path, "w", encoding="utf-8") as f:
        f.write(content[:m.start()] + content[m.end():])
    return "".join(LATEX_ESPECIALS.get(c, c) for c in m.group(1))


def elimina_instruccions(path):
    with open(path, encoding="utf-8") as f:
        content = f.read()
    with open(path, "w", encoding="utf-8") as f:
        f.write(BLOCKQUOTE_RE.sub("", content))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("ciclo")
    parser.add_argument("familia")
    parser.add_argument("--pd-dir", required=True)
    parser.add_argument("--stage", required=True)
    args = parser.parse_args()
    cicle = args.ciclo.upper()
    familia = args.familia.upper()

    with open(os.path.join(PROJECT_DIR, f"boe_{familia}", f"rd-{cicle.lower()}.json"), encoding="utf-8") as f:
        moduls = {str(k): v["nombre"] for k, v in json.load(f)["ModulosProfesionales"].items()}

    status = compute_pd_status(cicle, familia, args.pd_dir)
    draft_path = os.path.join(args.stage, ".draft")
    if is_pd_verified(status):
        if os.path.exists(draft_path):
            os.remove(draft_path)
    else:
        open(draft_path, "w").close()
        print(" * [PD] Report no verificat: s'aplica la marca d'aigua ESBORRANY")

    with open(os.path.join(args.stage, "pd_header.tex"), "w", encoding="utf-8") as f:
        f.write(HEADER_TEX)

    excel = status["excel_path"]
    sheetnames = []
    exportar = None
    if os.path.exists(excel):
        import openpyxl
        wb = openpyxl.load_workbook(excel, read_only=True)
        sheetnames = wb.sheetnames
        wb.close()
        exportar = load_excel_exporter()
    else:
        print(f" * [PD] Excel no trobat ({excel}): sense Quadres Resum")

    for fname in sorted(os.listdir(args.stage)):
        parsed = parse_pd_filename(fname)
        if not parsed:
            continue
        path = os.path.join(args.stage, fname)
        codi = parsed["codi"]
        nombre = moduls.get(codi)
        fulla = resol_fulla(sheetnames, nombre) if nombre else None

        excel_ko = bool(fulla) and any(f"Fulla '{fulla}'" in e for e in status["excel_issues"])
        marca = ""
        if parsed["estat"] == "BORRADOR":
            marca += MARCA_BORRADOR
        if find_placeholders(path) or excel_ko:
            marca += MARCA_INCIDENCIES
        if marca:
            marca_titol(path, marca)
        elimina_instruccions(path)

        if not (exportar and fulla):
            if exportar and nombre:
                print(f" * [PD] {codi}: fulla '{nombre}' no trobada a l'Excel, sense Quadre Resum")
            continue
        tmp_pdf = "/tmp/cuadro-resumen.pdf"
        if os.path.exists(tmp_pdf):
            os.remove(tmp_pdf)
        cwd = os.getcwd()
        os.chdir(PROJECT_DIR)  # exportar_rango_a_pdf usa ./temp/ relatiu
        try:
            exportar(excel, fulla, "B1:I200", tmp_pdf)
        finally:
            os.chdir(cwd)
        if not os.path.exists(tmp_pdf):
            print(f" * [PD] {codi}: no s'ha pogut exportar el Quadre Resum")
            continue
        pdf_name = f"PD_9999_{codi}_CuadroResumen.pdf"
        shutil.move(tmp_pdf, os.path.join(args.stage, pdf_name))
        titol = extrau_titol_esquema(path)
        toc = f",addtotoc={{1,subsection,2,{{{titol}}},esquema-{codi}}}" if titol else ""
        with open(path, "a", encoding="utf-8") as f:
            f.write(f"\n\n\\includepdf[pages=-,landscape=true{toc}]{{./{pdf_name}}}\n")
        print(f" * [PD] {codi}: Quadre Resum afegit")


if __name__ == "__main__":
    main()
